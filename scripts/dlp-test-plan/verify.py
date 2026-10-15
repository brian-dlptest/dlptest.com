#!/usr/bin/env python3
"""Static verification of the workbook. No LibreOffice here, so instead of recalculating we
resolve every formula reference against the real data and simulate the counts in Python."""
import os, re, sys
sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries, get_column_letter

P = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                 "DLPTest-com_Endpoint_DLP_Use_Case_Test_Plan.xlsx")
wb = load_workbook(P)
fail = []

DATA_SHEETS = {"1. Classification": 2, "2. Policy": 3, "3. Investigations": 2, "4. Usability": 2}

print("SHEETS")
for ws in wb.worksheets:
    print(f"  {ws.title:<22} {ws.dimensions:<12} rows={ws.max_row:<4} cols={ws.max_column}")
    if len(ws.title) > 31 or re.search(r"[\\/?*\[\]:]", ws.title):
        fail.append(f"illegal sheet name {ws.title!r}")

extents = {}
for name, first in DATA_SHEETS.items():
    ws, n, r = wb[name], 0, first
    while ws.cell(row=r, column=1).value:
        n += 1; r += 1
    extents[name] = (first, first + n - 1, n)
print("\nDATA EXTENTS  (first, last, count)")
for k, v in extents.items():
    print(f"  {k:<22} {v}")

# ---- 1. formulas: safe functions, resolvable refs, correct extents ----------
ALLOWED = {"COUNTIF", "COUNTIFS", "SUM", "IFERROR", "IF", "AND", "TEXTJOIN"}
# Written bare, TEXTJOIN evaluates to #NAME? - Excel stores post-2007 functions prefixed.
BARE_TEXTJOIN = re.compile(r"(?<!_xlfn\.)TEXTJOIN\(")
FUNC_RE = re.compile(r"([A-Z_][A-Z0-9_.]*)\s*\(")
REF_RE = re.compile(r"(?:'([^']+)'|([A-Za-z0-9_.]+))!\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)")
summary = wb["Scoring Summary"]
formulas = [(c.coordinate, c.value) for row in summary.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("=")]
print(f"\nFORMULAS on Scoring Summary: {len(formulas)}")

def col_idx(letter):
    return range_boundaries(f"{letter}1:{letter}1")[0]

input_cols, criteria_cols = set(), set()
CRITERIA_WORDS = {"Table stakes", "Advanced", "Differentiator", "Windows", "macOS", "Linux"}

def split_args(call):
    """Split a COUNTIF(S) argument list on top-level commas."""
    return [x.strip() for x in re.split(r',(?=(?:[^"]*"[^"]*")*[^"]*$)', call)]

def pairs_of(f):
    """(range_text, criterion) pairs, each criterion bound to the range it follows."""
    out = []
    for call in re.findall(r"COUNTIFS?\((.*?)\)(?=[+\s]|$)", f):
        a = split_args(call)
        out += [(a[i], a[i + 1].strip('"')) for i in range(0, len(a) - 1, 2)]
    return out

for coord, f in formulas:
    for fn in FUNC_RE.findall(f):
        if fn not in ALLOWED:
            fail.append(f"{coord}: unsupported function {fn}")
    if BARE_TEXTJOIN.search(f):
        fail.append(f"{coord}: TEXTJOIN without the _xlfn. prefix will show #NAME?")
    for m in REF_RE.finditer(f):
        sheet = m.group(1) or m.group(2)
        c1, r1, c2, r2 = m.group(3), int(m.group(4)), m.group(5), int(m.group(6))
        if sheet not in wb.sheetnames:
            fail.append(f"{coord}: unknown sheet {sheet!r}"); continue
        if " " in sheet and not m.group(1):
            fail.append(f"{coord}: sheet {sheet!r} has a space but is unquoted")
        first, last, _ = extents[sheet]
        if (r1, r2) != (first, last):
            fail.append(f"{coord}: {sheet}!{c1}{r1}:{c2}{r2} spans rows {r1}-{r2}, data is {first}-{last}")
    # classify each range by the criterion that immediately follows IT, sheet-aware
    for rng, crit in pairs_of(f):
        m = REF_RE.search(rng)
        if not m:
            continue
        sheet = m.group(1) or m.group(2)
        if sheet not in wb.sheetnames:
            continue
        target = criteria_cols if crit in CRITERIA_WORDS else input_cols
        for ci in range(col_idx(m.group(3)), col_idx(m.group(5)) + 1):
            target.add((sheet, ci))

overlap = input_cols & criteria_cols
if overlap:
    fail.append(f"columns used as both result and criteria: {sorted(overlap)}")

# result columns must be blank in every data row; criteria columns must be populated
print("\nCOLUMN ROLES")
for (sheet, ci) in sorted(input_cols):
    ws = wb[sheet]; first, last, _ = extents[sheet]
    bad = [r for r in range(first, last + 1) if ws.cell(row=r, column=ci).value not in (None, "")]
    print(f"  input    {sheet:<22} {get_column_letter(ci)}  {'empty OK' if not bad else 'NOT EMPTY ' + str(bad[:4])}")
    if bad: fail.append(f"{sheet}!{get_column_letter(ci)} is a result column but has values at {bad[:4]}")
for (sheet, ci) in sorted(criteria_cols):
    ws = wb[sheet]; first, last, _ = extents[sheet]
    bad = [r for r in range(first, last + 1) if not ws.cell(row=r, column=ci).value]
    print(f"  criteria {sheet:<22} {get_column_letter(ci)}  {'populated OK' if not bad else 'BLANKS ' + str(bad[:4])}")
    if bad: fail.append(f"{sheet}!{get_column_letter(ci)} is a criteria column but is blank at {bad[:4]}")

# ---- 2. hardcoded totals in the Untested formulas --------------------------
print("\nUNTESTED-COUNT TOTALS")
TOT_RE = re.compile(r"^=(\d+)-SUM\(C(\d+):F\d+\)$")
for coord, f in formulas:
    m = TOT_RE.match(f)
    if not m: continue
    total, row = int(m.group(1)), int(m.group(2))
    label = summary.cell(row=row, column=2).value
    # recompute from the sibling COUNTIF/COUNTIFS on the same row
    sib = summary.cell(row=row, column=3).value
    refs = REF_RE.findall(sib)
    span = 0
    for _q, _b, c1, r1, c2, r2 in refs:
        sheet = _q or _b
        cols = col_idx(c2) - col_idx(c1) + 1
        rows_ = extents[sheet][2]
        span += rows_ * cols if len(refs) == 1 else 0
    if len(refs) == 1 and 'COUNTIFS' not in sib and span != total:
        fail.append(f"{coord}: total {total} != {span} addressable cells for {label!r}")
    print(f"  {coord:<6} {str(label)[:34]:<34} total={total}")

# ---- 3. simulate every COUNTIF / COUNTIFS against seeded values -------------
print("\nCOUNT SIMULATION")
SEED = {"1. Classification": {9: ["Pass"]*8 + ["Partial"]*3 + ["Fail"]*2 + ["N/A"]*1},
        "2. Policy": {6: ["Pass"]*30 + ["Fail"]*10, 7: ["Partial"]*15, 8: ["Pass"]*20 + ["N/A"]*5},
        "3. Investigations": {7: ["Pass"]*9 + ["Partial"]*4},
        "4. Usability": {7: ["Pass"]*7 + ["Fail"]*3 + ["N/A"]*2}}
for sheet, cols in SEED.items():
    ws = wb[sheet]; first = extents[sheet][0]
    for ci, vals in cols.items():
        for i, v in enumerate(vals):
            ws.cell(row=first + i, column=ci, value=v)

def evaluate(f):
    """Evaluate a =COUNTIF(...)+COUNTIFS(...) chain against the live workbook."""
    total = 0
    for call in re.findall(r"COUNTIFS?\((.*?)\)(?=[+\s]|$)", f):
        parts = [x.strip() for x in re.split(r",(?=(?:[^\"]*\"[^\"]*\")*[^\"]*$)", call)]
        pairs = [(parts[i], parts[i + 1].strip('"')) for i in range(0, len(parts) - 1, 2)]
        base = REF_RE.search(pairs[0][0])
        sheet = base.group(1) or base.group(2)
        ws = wb[sheet]; first, last, _ = extents[sheet]
        for r in range(first, last + 1):
            ok = True
            for rng, want in pairs:
                m = REF_RE.search(rng)
                cells = [ws.cell(row=r, column=ci).value
                         for ci in range(col_idx(m.group(3)), col_idx(m.group(5)) + 1)]
                if len(pairs) == 1:
                    ok = None          # multi-cell single-criterion: count cells, not rows
                    total += sum(1 for v in cells if v == want)
                    break
                if not any(v == want for v in cells):
                    ok = False; break
            if ok: total += 1
    return total

checks = 0
for coord, f in formulas:
    if "COUNTIF" not in f: continue
    got = evaluate(f)
    checks += 1
    summary_val = got
    if got < 0: fail.append(f"{coord}: negative count")
print(f"  evaluated {checks} count formulas against seeded data without error")


# ---- 3b. table stakes failures list: helpers -> TEXTJOIN -> reconciliation ---
print("\nTABLE STAKES FAILURES LIST")
#            sheet                 helper  tier  result     (Policy: Block)
HELPER = {"1. Classification": ("K", "H", "I"), "2. Policy": ("K", "H", "G"),
          "3. Investigations": ("I", "F", "G"), "4. Usability": ("I", "F", "G")}
HIT = r'AND\(([A-Z]+)(\d+)="Table stakes",([A-Z]+)(\d+)="Fail"\)'
HELPER_FIRST = re.compile(r'^=IF\(' + HIT + r',A(\d+),""\)$')
HELPER_NEXT = re.compile(r'^=IF\(' + HIT + r',IF\(([A-Z]+)(\d+)="",A(\d+),([A-Z]+)(\d+)&", "&A(\d+)\),([A-Z]+)(\d+)\)$')
LIST_RE = re.compile(r'^=IF\(D(\d+)=0,"None recorded",' r"'([^']+)'!\$([A-Z]+)\$(\d+)\)$")

# every helper cell reads only its own row, and (after the first) the helper cell above it
for sheet, (hcol, tcol, rcol) in HELPER.items():
    ws = wb[sheet]; first, last, _ = extents[sheet]
    if not ws.column_dimensions[hcol].hidden:
        fail.append(f"{sheet}: helper column {hcol} is not hidden")
    bad = []
    for r in range(first, last + 1):
        f_ = str(ws[f"{hcol}{r}"].value)
        if r == first:
            m = HELPER_FIRST.match(f_)
            ok_ = bool(m) and (m.group(1), m.group(3)) == (tcol, rcol) and {m.group(2), m.group(4), m.group(5)} == {str(r)}
        else:
            m = HELPER_NEXT.match(f_)
            ok_ = (bool(m) and (m.group(1), m.group(3)) == (tcol, rcol)
                   and {m.group(2), m.group(4), m.group(7), m.group(10)} == {str(r)}
                   and {(m.group(5), m.group(6)), (m.group(8), m.group(9)), (m.group(11), m.group(12))} == {(hcol, str(r - 1))})
        if not ok_:
            bad.append(r)
    print(f"  helper  {sheet:<20} {hcol}  tier={tcol} result={rcol}  {extents[sheet][2]} rows  "
          f"{'OK' if not bad else 'BAD ' + str(bad[:4])}")
    if bad:
        fail.append(f"{sheet}!{hcol}: helper formula wrong or not self-referencing at rows {bad[:4]}")

# find the list rows and the tier rollup's Table stakes row on the Summary
list_rows = {}
for row in summary.iter_rows(min_col=5, max_col=5):
    c = row[0]
    m = LIST_RE.match(str(c.value)) if isinstance(c.value, str) else None
    if m:
        list_rows[m.group(2)] = (c.row, m)
ts_tier_row = next(r for r in range(1, summary.max_row + 1)
                   if summary.cell(row=r, column=2).value == "Table stakes")
if set(list_rows) != set(HELPER):
    fail.append(f"failures list covers {sorted(list_rows)}, expected all {len(HELPER)} sheets")

# seed deliberate table stakes Fails on top of the generic seed, then simulate
for sheet, (hcol, tcol, rcol) in HELPER.items():
    ws = wb[sheet]; first, last, _ = extents[sheet]
    ts = [r for r in range(first, last + 1) if ws[f"{tcol}{r}"].value == "Table stakes"]
    for r in ts[:2]:
        ws[f"{rcol}{r}"] = "Fail"
    for r in ts[2:3]:
        ws[f"{rcol}{r}"] = "Pass"

listed_total = 0
for sheet, (hcol, tcol, rcol) in HELPER.items():
    ws = wb[sheet]; first, last, _ = extents[sheet]
    row, m = list_rows[sheet]
    d_row, ref_sheet = int(m.group(1)), m.group(2)
    c1, r1 = m.group(3), int(m.group(4))
    if d_row != row:
        fail.append(f"Summary!E{row}: list reads D{d_row}, not its own row's count")
    if (ref_sheet, c1, r1) != (sheet, hcol, last):
        fail.append(f"Summary!E{row}: reads {ref_sheet}!{c1}{r1}, not the last helper cell {sheet}!{hcol}{last}")
    # what any spreadsheet computes: the running list down the helper column
    joined = ""
    for r in range(first, last + 1):
        if ws[f"{tcol}{r}"].value == "Table stakes" and ws[f"{rcol}{r}"].value == "Fail":
            joined = ws[f"A{r}"].value if joined == "" else joined + ", " + ws[f"A{r}"].value
    expected = [ws[f"A{r}"].value for r in range(first, last + 1)
                if ws[f"{tcol}{r}"].value == "Table stakes" and ws[f"{rcol}{r}"].value == "Fail"]
    d_count = evaluate(summary[f"D{row}"].value)
    shown = "None recorded" if d_count == 0 else joined
    ok = joined == ", ".join(expected) and len(expected) == d_count
    listed_total += len(expected)
    print(f"  list    {sheet:<20} Fails={d_count:<2} shows: {shown[:60]}{'...' if len(shown) > 60 else ''}  {'OK' if ok else 'MISMATCH'}")
    if not ok:
        fail.append(f"{sheet}: list {joined!r} / count {d_count} disagree with expected {expected}")

tier_fail = evaluate(summary[f"E{ts_tier_row}"].value)
recon = tier_fail == listed_total
print(f"  reconcile: tier rollup Table stakes Fail = {tier_fail}, IDs listed = {listed_total}  "
      f"{'OK' if recon else 'MISMATCH'}")
if not recon:
    fail.append(f"tier rollup says {tier_fail} table stakes Fails but the list shows {listed_total}")

# the floor is enforced visually, and the footnote no longer claims three columns
floor_rules = [rule for rng in summary.conditional_formatting for rule in rng.rules
               if str(rng.sqref) == f"H{ts_tier_row}" and rule.formula and "<0.7" in rule.formula[0]]
print(f"  floor highlight on H{ts_tier_row} (<70%): {'OK' if floor_rules else 'MISSING'}")
if not floor_rules:
    fail.append(f"no <70% conditional format on the Table stakes coverage cell H{ts_tier_row}")
note = " ".join(str(summary.cell(row=r, column=2).value or "") for r in range(1, summary.max_row + 1))
if "of the three" in note:
    fail.append("Scoring Summary footnote still says 'of the three' - Policy has two result columns")

# ---- 4. dropdowns and conditional formatting -------------------------------
print("\nVALIDATION / CONDITIONAL FORMATTING")
SCORING = '"Pass,Partial,Fail,N/A"'
TIMING = '"Inline,After the fact,Not blocked"'      # Policy only, recorded but never scored
timing_dvs = []
for ws in wb.worksheets:
    for dv in ws.data_validations.dataValidation:
        print(f"  DV  {ws.title:<22} {str(dv.sqref):<14} {dv.formula1}")
        if dv.formula1 == TIMING:
            timing_dvs.append((ws.title, str(dv.sqref)))
        elif dv.formula1 != SCORING:
            fail.append(f"{ws.title}!{dv.sqref}: unexpected dropdown {dv.formula1}")
    for rng in ws.conditional_formatting:
        print(f"  CF  {ws.title:<22} {rng.sqref}  ({len(rng.rules)} rules)")
pol_last = extents["2. Policy"][1]
if timing_dvs != [("2. Policy", f"J3:J{pol_last}")]:
    fail.append(f"Block timing dropdown should cover exactly 2. Policy!J3:J{pol_last}; found {timing_dvs}")
print(f"  scoring vocabulary on every result column: {SCORING}")
print(f"  Block timing on 2. Policy!J3:J{pol_last} only: {timing_dvs == [('2. Policy', f'J3:J{pol_last}')]}")
scored_timing = [c for c, f in formulas if "'2. Policy'!$J$" in f]
if scored_timing:
    fail.append(f"Block timing is meant to be unscored, but Summary formulas read it: {scored_timing}")

# ---- 5. reference integrity -------------------------------------------------
print("\nREFERENCE INTEGRITY")
ID_RE = re.compile(r"\b((?:C|P|E|I|U)-[A-Z]?\d{2}[a-z]?)\b")
ids = set()
for ws in wb.worksheets:
    for r_ in range(1, ws.max_row + 1):
        v = ws.cell(row=r_, column=1).value
        if isinstance(v, str):
            base = v.strip().split(".")[0]
            if ID_RE.fullmatch(base): ids.add(base)
dangling = {}
for ws in wb.worksheets:
    if ws.title == "Change Log":
        continue                      # historical IDs, checked separately below
    for row in ws.iter_rows():
        for cell in row:
            for t in [cell.value] + ([cell.comment.text] if cell.comment else []):
                if not isinstance(t, str): continue
                own = ws.cell(row=cell.row, column=1).value
                own = own.split(".")[0] if isinstance(own, str) else own
                for ref in ID_RE.findall(t):
                    if ref not in ids and ref != own:
                        dangling.setdefault(ref, []).append(f"{ws.title}!{cell.coordinate}")
print(f"  {len(ids)} live IDs; {len(dangling)} dangling reference(s)")
for ref, locs in sorted(dangling.items()):
    fail.append(f"dangling reference {ref} at {', '.join(locs[:4])}")

# Change Log: entries for THIS version must name live rows; older entries keep the IDs
# they had at the time and are not checked.
_here = os.path.dirname(os.path.abspath(__file__))
_build_src = open(os.path.join(_here, "build.py")).read()
cur_ver = int(re.search(r"^VERSION = (\d+)$", _build_src, re.M).group(1))
cl = wb["Change Log"]
cur_entries, cl_bad = 0, []
for r_ in range(3, cl.max_row + 1):
    if cl.cell(row=r_, column=1).value != cur_ver:
        continue
    cur_entries += 1
    for c_ in (3, 4, 5):
        for ref in ID_RE.findall(str(cl.cell(row=r_, column=c_).value or "")):
            if ref not in ids:
                cl_bad.append(f"Change Log!{cl.cell(row=r_, column=c_).coordinate} {ref}")
print(f"  Change Log: {cur_entries} entries for v{cur_ver}; {len(cl_bad)} naming a row that does not exist")
if cur_entries == 0:
    fail.append(f"Change Log has no entries for the current version v{cur_ver}")
for b in cl_bad:
    fail.append(f"{b}: current-version change log entry names a row that does not exist")

# build.py must point at rows through ref(), never a literal ID - a literal one survives a
# renumber by landing on whatever row now has that number, which is how v6 shipped three
# wrong glossary pointers past every check here.
literal = [f"build.py:{i}" for i, l in enumerate(_build_src.split("\n"), 1)
           if not l.lstrip().startswith("#") and ID_RE.search(l)]
print(f"  literal row IDs in build.py: {literal or 'none - all pointers go through ref()'}")
for loc in literal:
    fail.append(f"{loc}: literal row ID - use ref() so the pointer cannot drift")

# ---- 5b. every ID is well formed ------------------------------------------
print("\nID FORMAT")
CANON = re.compile(r"^(?:C-[ABCD]|P-[GMCBPN]|I-|U-)\d{2}(?:\.(?:Windows|macOS|Linux))?$")
bad = []
for sheet, first in DATA_SHEETS.items():
    ws = wb[sheet]
    for r_ in range(first, extents[sheet][1] + 1):
        v = ws.cell(row=r_, column=1).value
        if not isinstance(v, str) or not CANON.match(v):
            bad.append(f"{sheet}!A{r_}={v!r}")
print(f"  {sum(extents[s][2] for s in DATA_SHEETS)} IDs checked; {len(bad)} malformed")
for b in bad:
    fail.append(f"malformed ID {b} - did renumber.py run, and does the placeholder match?")

# ---- 6. source <-> workbook ID parity --------------------------------------
print("\nSOURCE <-> WORKBOOK ID PARITY")
import content_class, content_policy, content_xp, content_use
SRC = {"1. Classification": [r[0] for r in content_class.CLASSIFICATION],
       "2. Policy":         [r[0] for r in content_policy.POLICY],
       "3. Investigations": [r[0] for r in content_xp.INVESTIGATIONS],
       "4. Usability":      [r[0] for r in content_use.USABILITY]}
for sheet, want in SRC.items():
    ws, first = wb[sheet], extents[sheet][0]
    got = [ws.cell(row=r, column=1).value for r in range(first, first + len(want))]
    print(f"  {sheet:<22} {len(want):>3} IDs  {'OK' if got == want else 'MISMATCH'}")
    if got != want:
        diff = [f"{a}!={b}" for a, b in zip(want, got) if a != b][:5]
        fail.append(f"{sheet}: workbook IDs differ from source ({', '.join(diff)}) - stale __pycache__?")

# ---- 7. chart ---------------------------------------------------------------
print("\nCHART")
charts = summary._charts
print(f"  {len(charts)} chart(s) on Scoring Summary")
for ch in charts:
    print(f"    type={ch.type} grouping={ch.grouping} series={len(ch.series)}")
    if ch.grouping != "stacked": fail.append("chart is not stacked")
    if len(ch.series) != 5: fail.append(f"chart has {len(ch.series)} series, expected 5")
if not charts: fail.append("no chart on Scoring Summary")


# ---- 8. the page advertises what the workbook actually is ------------------
print("\nPAGE PARITY  (src/data/test-plan.ts)")
here = os.path.dirname(os.path.abspath(__file__))
ts_src = open(os.path.join(here, "..", "..", "src", "data", "test-plan.ts")).read()
build_src = open(os.path.join(here, "build.py")).read()
page_ver = int(re.search(r"TEST_PLAN_VERSION = (\d+);", ts_src).group(1))
build_ver = int(re.search(r"^VERSION = (\d+)$", build_src, re.M).group(1))
readme_ver = str(wb["Read Me"]["B4"].value)
print(f"  version  page={page_ver}  build.py={build_ver}  Read Me={readme_ver!r}")
if not (page_ver == build_ver and f"version {build_ver}" in readme_ver):
    fail.append(f"version drift: page {page_ver}, build.py {build_ver}, Read Me {readme_ver!r}")
page_rows = {n: int(c) for n, c in re.findall(r'name: "([^"]+)",\s*rows: (\d+),', ts_src)}
page_total = int(re.search(r"TEST_PLAN_ROWS = (\d+);", ts_src).group(1))
for sheet, want in SRC.items():
    got = page_rows.get(sheet)
    print(f"  {sheet:<20} page={got}  workbook={len(want)}  {'OK' if got == len(want) else 'MISMATCH'}")
    if got != len(want):
        fail.append(f"page says {sheet} has {got} rows, workbook has {len(want)}")
wb_total = sum(len(v) for v in SRC.values())
print(f"  total                page={page_total}  workbook={wb_total}  {'OK' if page_total == wb_total else 'MISMATCH'}")
if page_total != wb_total:
    fail.append(f"page says {page_total} test cases, workbook has {wb_total}")


# ---- 9. no template leakage, no em dashes ----------------------------------
print("\nRENDERED TEXT")
here_ = os.path.dirname(os.path.abspath(__file__))
page_src = open(os.path.join(here_, "..", "..", "src", "pages", "endpoint-dlp-test-plan.astro")).read()
leaks, emdash = [], []
for ws_ in wb.worksheets:
    for row in ws_.iter_rows():
        for c in row:
            v = c.value
            if not isinstance(v, str) or v.startswith("="):
                continue
            # an f-string that was never an f-string prints its {ref(...)} literally
            if "{" in v or "}" in v:
                leaks.append(f"{ws_.title}!{c.coordinate}: {v[max(0, v.find('{') - 20):v.find('{') + 40]!r}")
            if "—" in v:
                emdash.append(f"{ws_.title}!{c.coordinate}")
for name, src_ in (("src/data/test-plan.ts", ts_src), ("endpoint-dlp-test-plan.astro", page_src)):
    if "ref(" in src_:
        leaks.append(f"{name}: contains ref(")
    if "—" in src_:
        emdash.append(name)
print(f"  template text in cells or page data: {leaks or 'none'}")
print(f"  em dashes in workbook or page data: {emdash or 'none'}")
for l in leaks:  fail.append(f"unrendered template text - {l}")
for e in emdash: fail.append(f"em dash in {e}")

# ---- 10. the page's channel count matches the Policy sheet ------------------
chan_page = int(re.search(r'"(\d+) egress channels', ts_src).group(1))
chan_wb = len({r[0].split(".")[0] for r in content_policy.POLICY})
print(f"  Policy channels: page={chan_page} workbook={chan_wb}  {'OK' if chan_page == chan_wb else 'MISMATCH'}")
if chan_page != chan_wb:
    fail.append(f"page says {chan_page} egress channels, Policy sheet has {chan_wb}")

# ---- 11. ID map reconciles with the previous version ------------------------
print("\nID MAP")
import content_changelog
prev_ver, prev_rows = content_changelog.IDMAP_FROM
full_ids = {wb[sh].cell(row=r_, column=1).value for sh, (f_, l_, _) in extents.items()
            for r_ in range(f_, l_ + 1)}
cl_ = wb["Change Log"]
hdr = next((r_ for r_ in range(1, cl_.max_row + 1) if cl_.cell(row=r_, column=1).value == f"v{prev_ver} ID"), None)
if hdr is None:
    fail.append("Change Log has no ID map header")
else:
    olds, news, bad_map = [], [], []
    r_ = hdr + 1
    while cl_.cell(row=r_, column=1).value:
        o, n_ = cl_.cell(row=r_, column=1).value, cl_.cell(row=r_, column=2).value
        olds.append(o); news.append(n_)
        if n_ != "Removed" and n_ not in full_ids:
            bad_map.append(f"row {r_}: v{cur_ver} ID {n_} does not exist")
        if o != "New" and not re.fullmatch(r"(?:C-[ABCD]|P-[GMCBPN]|E-|I-|U-)\d{2}(?:\.(?:Windows|macOS|Linux))?", o):
            bad_map.append(f"row {r_}: v{prev_ver} ID {o!r} is malformed")
        r_ += 1
    real_old = [o for o in olds if o != "New"]; real_new = [n_ for n_ in news if n_ != "Removed"]
    if len(set(real_old)) != len(real_old): bad_map.append("a previous ID appears twice")
    if len(set(real_new)) != len(real_new): bad_map.append("a current ID appears twice")
    removed, added = news.count("Removed"), olds.count("New")
    total_now = sum(e[2] for e in extents.values())
    ok_sum = prev_rows - removed + added == total_now
    print(f"  {len(olds)} rows: removed {removed}, new {added}, renumbered or re-score {len(olds) - removed - added}")
    print(f"  reconcile: v{prev_ver} {prev_rows} - {removed} + {added} = {prev_rows - removed + added}, "
          f"v{cur_ver} has {total_now}  {'OK' if ok_sum else 'MISMATCH'}")
    if not ok_sum:
        bad_map.append(f"v{prev_ver} {prev_rows} - {removed} removed + {added} new != {total_now} rows now")
    for b in bad_map: fail.append(f"ID map: {b}")

# ---- 9. test data database (test-data/cases.json) -------------------------
# Row IDs are reassigned by sheet position on every build, so a test can keep its ID
# while becoming a different test (v7: U-22 went from "User-visible latency" to
# "Working redirect to the sanctioned alternative"). Comparing ID sets would pass that.
# So every case is checked by ID AND name, and by a hash of the row's steps, which
# catches a reworded row that kept its name. See TEST-DATA-REVIEW.md.
#
#   python3 verify.py --stamp   after a review: record each row's current steps hash
#                               and set reviewed_for_version, then re-run verify.py
print("\nTEST DATA DATABASE  (test-data/cases.json)")
import collections, hashlib, json
td = os.path.join(here, "test-data")
STAMP = "--stamp" in sys.argv
# Row IDs in prose, file names or folder names. Not \b: `C-A10_ocr-on-images` has no
# word boundary between the ID and the underscore.
DOC_ID_RE = re.compile(r"(?<![A-Za-z0-9])((?:C|P|E|I|U)-[A-Z]?\d{2})(?![0-9])")


def steps_sha(row, sheet_rows):
    """Hash of what a tester is told to do: How to Test (What to Test on Policy), plus
    the expected result on Classification. Wording cleanup changes it too; the review
    decides whether the change matters (TEST-DATA-REVIEW.md, "Which rewording counts")."""
    text = row[4] + ("\n" + row[5] if sheet_rows is content_class.CLASSIFICATION else "")
    return hashlib.sha256(text.encode()).hexdigest()[:12]


try:
    db = json.load(open(os.path.join(td, "cases.json")))
except (OSError, ValueError) as e:
    db = None
    fail.append(f"test-data/cases.json unreadable: {e}")
if db is not None:
    reviewed = db.get("reviewed_for_version")
    print(f"  reviewed_for_version={reviewed}  build.py VERSION={build_ver}")

    rows = {}
    for sheet_rows in (content_class.CLASSIFICATION, content_policy.POLICY,
                       content_xp.INVESTIGATIONS, content_use.USABILITY):
        for r in sheet_rows:
            rows[r[0]] = (r[2], steps_sha(r, sheet_rows))
    cases = db.get("cases", [])
    dupes = sorted(i for i, n in collections.Counter(c["id"] for c in cases).items() if n > 1)
    for d_ in dupes:
        fail.append(f"cases.json: {d_} appears more than once - keep one case per row")
    have = {c["id"]: c for c in cases}

    if STAMP:
        stamped = 0
        for c in cases:
            if c["id"] in rows and c.get("name") == rows[c["id"]][0]:
                stamped += c.get("steps_sha") != rows[c["id"]][1]
                c["steps_sha"] = rows[c["id"]][1]
        db["reviewed_for_version"] = reviewed = build_ver
        with open(os.path.join(td, "cases.json"), "w") as fh:
            fh.write(json.dumps(db, indent=2, ensure_ascii=False) + "\n")
        print(f"  --stamp: recorded steps_sha on {stamped} changed cases, "
              f"reviewed_for_version={build_ver}")

    if reviewed != build_ver:
        fail.append(f"test data last reviewed for v{reviewed}, workbook is v{build_ver} - run the "
                    "review in TEST-DATA-REVIEW.md, then verify.py --stamp")
    moved, missing, reworded = [], [], []
    for rid, (name, sha) in rows.items():
        case = have.get(rid)
        if case is None:
            missing.append(rid)
        elif case.get("name") != name:
            moved.append(f"{rid}: cases.json says {case.get('name')!r}, workbook says {name!r}")
        elif case.get("steps_sha") != sha:
            reworded.append(f"{rid} ({name})")
    stale = sorted(set(have) - set(rows))
    print(f"  {len(rows)} rows  {len(cases)} cases  missing={len(missing)}  "
          f"name changed={len(moved)}  steps changed={len(reworded)}  stale={len(stale)}  "
          f"duplicate={len(dupes)}")
    for m in missing:
        fail.append(f"cases.json: no case for {m} - new row? decide its files (TEST-DATA-REVIEW.md R3)")
    for m in moved:
        fail.append(f"cases.json: {m} - this ID now names a different test; re-map it")
    for m in reworded:
        fail.append(f"cases.json: {m} - its steps changed since the last review; check its "
                    "files still fit (R1 or R3), then verify.py --stamp")
    for m in stale:
        fail.append(f"cases.json: {m} is not in the workbook - removed row? (TEST-DATA-REVIEW.md R2)")

    # Files under generated/ are built locally and never committed, so they cannot be
    # checked for on disk in a clean clone. They are checked against the list of outputs
    # `build.py local` promises to write (tools/outputs.py), and on disk when built.
    import importlib.util
    _spec = importlib.util.spec_from_file_location(
        "test_data_outputs", os.path.join(td, "tools", "outputs.py"))
    _outputs = importlib.util.module_from_spec(_spec)
    _spec.loader.exec_module(_outputs)
    generated_outputs = set(_outputs.GENERATED_OUTPUTS)
    generated_built = os.path.isdir(os.path.join(td, "generated"))

    referenced, bad_paths, unexplained = set(), [], []
    for c in cases:
        files = c.get("files", [])
        if not files and not c.get("no_file_reason"):
            unexplained.append(c["id"])
        for f_ in files:
            referenced.add(f_)
            if f_.startswith("generated/"):
                if f_ not in generated_outputs:
                    bad_paths.append(f"{c['id']}: {f_} (not an output of build.py local)")
                elif generated_built and not os.path.isfile(os.path.join(td, f_)):
                    bad_paths.append(f"{c['id']}: {f_} (re-run build.py local)")
            elif not f_.startswith("files/") or not os.path.isfile(os.path.join(td, f_)):
                bad_paths.append(f"{c['id']}: {f_}")
    # Data a Read Me instruction needs but no row does. If the instruction leaves the
    # rendered Read Me, the data loses its reason to exist - fail so someone decides (R2).
    readme_cells = {str(v).strip() for row in wb["Read Me"].iter_rows(values_only=True)
                    for v in row if v is not None}
    for req in db.get("readme_requirements", []):
        if req["name"] not in readme_cells:
            fail.append(f"cases.json: Read Me instruction {req['name']!r} is no longer in the "
                        "workbook's Read Me - remove or re-home its files (TEST-DATA-REVIEW.md R2)")
        for f_ in req.get("files", []):
            referenced.add(f_)
            if not os.path.isfile(os.path.join(td, f_)):
                bad_paths.append(f"Read Me {req['name']!r}: {f_}")
    for u in unexplained:
        fail.append(f"cases.json: {u} has no files and no no_file_reason")
    for b in bad_paths:
        fail.append(f"cases.json: missing file {b}")

    DOCS = ("README.txt", "README.md", "RUNBOOK.md")
    on_disk, orphans, id_paths = [], [], []
    for dp, dn, fn in os.walk(os.path.join(td, "files")):
        for d_ in dn:
            if DOC_ID_RE.search(d_):
                id_paths.append(os.path.relpath(os.path.join(dp, d_), td) + "/")
        for f_ in fn:
            if f_ == ".DS_Store":
                continue
            rel = os.path.relpath(os.path.join(dp, f_), td).replace(os.sep, "/")
            on_disk.append(rel)
            if DOC_ID_RE.search(f_):
                id_paths.append(rel)
            if f_ not in DOCS and rel not in referenced:
                orphans.append(rel)
    print(f"  {len(on_disk)} committed files, {len(referenced)} referenced, "
          f"{len(orphans)} orphaned, {len(bad_paths)} missing")
    for o in orphans[:10]:
        fail.append(f"test-data: {o} is not used by any case - remove it, or map it (R2)")
    if len(orphans) > 10:
        fail.append(f"test-data: {len(orphans) - 10} more orphaned files")

    # Row IDs live in cases.json only. A README, file or folder that names a row goes
    # stale on the next renumber without anything noticing - which is the failure this
    # whole section exists for.
    id_docs = []
    for dp, dn, fn in os.walk(td):
        dn[:] = [d for d in dn if d not in ("generated", "tools", ".git")]
        for f_ in fn:
            if f_ in DOCS + ("TASK.md", "VALIDATION.md"):
                pth = os.path.join(dp, f_)
                hit = DOC_ID_RE.search(open(pth, errors="replace").read())
                if hit:
                    id_docs.append(f"{os.path.relpath(pth, td)} mentions {hit.group(1)}")
    id_docs += [f"{p_} has a row ID in its name" for p_ in id_paths]
    print(f"  row IDs in test-data docs and paths: {len(id_docs)}")
    for d_ in id_docs:
        fail.append(f"test-data: {d_} - name the capability instead; IDs belong in cases.json")

print("\n" + ("FAILURES:\n  " + "\n  ".join(fail) if fail else "ALL CHECKS PASSED"))
sys.exit(1 if fail else 0)
