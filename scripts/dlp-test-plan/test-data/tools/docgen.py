"""Rendering helpers: CSV/TXT, Word, Excel, PowerPoint, images, archives.

Paths are relative to test-data/. Committed payloads go under files/; anything that
cannot be committed (credential-shaped files, nested git repositories) goes under the
gitignored generated/ folder.
"""
import csv
import os
import shutil
import subprocess
import tarfile
import tempfile
import zipfile

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt as PPt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Containers (Office files, zip, tar.gz, 7z) embed the time they were written. Every
# entry is stamped with this instead, so a refresh replayed from the MCP cache is
# byte-identical and produces no diff. Only the password-protected zips still change:
# ZipCrypto salts each archive with random bytes.
FIXED_TIME = (2026, 1, 1, 0, 0, 0)
FIXED_EPOCH = 1767225600  # 2026-01-01T00:00:00Z


def _pin_zip_times(p):
    """Rewrite a zip-based file (docx/xlsx/pptx/zip) with every entry at FIXED_TIME."""
    import re
    with zipfile.ZipFile(p) as z:
        entries = [(i, z.read(i)) for i in z.infolist()]
    stamp = "%04d-%02d-%02dT%02d:%02d:%02dZ" % FIXED_TIME
    for n, (info, data) in enumerate(entries):
        if info.filename == "docProps/core.xml":
            # openpyxl stamps `modified` with the save time whatever the caller set.
            data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*",
                          rb"\g<1>" + stamp.encode(), data)
            entries[n] = (info, data)
    with zipfile.ZipFile(p, "w") as z:
        for info, data in entries:
            pinned = zipfile.ZipInfo(info.filename, date_time=FIXED_TIME)
            pinned.compress_type = info.compress_type
            pinned.external_attr = info.external_attr
            z.writestr(pinned, data)


def _fixed_datetime():
    import datetime
    return datetime.datetime(*FIXED_TIME)


def syn(*parts):
    """Join fragments at runtime so a credential-shaped literal never appears whole in
    source. GitHub push protection matches on shape, not validity - the same reason
    src/lib/regex/dlp-patterns.ts has a syn() helper."""
    return "".join(parts)


def path(*parts):
    """Resolve a library-relative path, creating parent directories."""
    p = os.path.join(ROOT, *parts)
    os.makedirs(os.path.dirname(p), exist_ok=True)
    return p


def replace_dir(rel, write):
    """Rebuild a whole folder without a window where it is missing: `write(tmp_rel)`
    fills a sibling folder, which then replaces `rel`. Fetch the data before calling
    this, so a network failure leaves the committed folder as it was."""
    final = os.path.join(ROOT, *rel.split("/"))
    tmp_rel = rel + ".partial"
    tmp = os.path.join(ROOT, *tmp_rel.split("/"))
    shutil.rmtree(tmp, ignore_errors=True)
    write(tmp_rel)
    shutil.rmtree(final, ignore_errors=True)
    os.rename(tmp, final)


def write_text(rel, text):
    p = path(*rel.split("/"))
    with open(p, "w") as fh:
        fh.write(text if text.endswith("\n") else text + "\n")
    return p


def write_csv(rel, header, rows):
    p = path(*rel.split("/"))
    with open(p, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        w.writerows(rows)
    return p


def write_docx(rel, title, paragraphs=None, table=None, heading_level=1):
    """Word document: a title, optional paragraphs, and one optional table."""
    doc = Document()
    doc.add_heading(title, heading_level)
    for para in paragraphs or []:
        if isinstance(para, tuple):
            kind, text = para
            if kind == "h":
                doc.add_heading(text, 2)
                continue
            if kind == "b":
                p = doc.add_paragraph()
                p.add_run(text).bold = True
                continue
        doc.add_paragraph(para)
    if table:
        header, rows = table
        t = doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        for i, h in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = str(h)
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
    doc.core_properties.created = doc.core_properties.modified = _fixed_datetime()
    p = path(*rel.split("/"))
    doc.save(p)
    _pin_zip_times(p)
    return p


def write_xlsx(rel, sheets):
    """Excel workbook. `sheets` is [(name, header, rows), ...]."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, header, rows in sheets:
        ws = wb.create_sheet(title=name[:31])
        if header:
            ws.append(list(header))
            for c in ws[1]:
                c.font = c.font.copy(bold=True)
        for row in rows:
            ws.append(list(row))
    wb.properties.created = wb.properties.modified = _fixed_datetime()
    p = path(*rel.split("/"))
    wb.save(p)
    _pin_zip_times(p)
    return p


def write_pptx(rel, slides):
    """Deck. `slides` is [(title, [bullet, ...]), ...]."""
    prs = Presentation()
    title_layout, body_layout = prs.slide_layouts[0], prs.slide_layouts[1]
    first = True
    for title, bullets in slides:
        slide = prs.slides.add_slide(title_layout if first else body_layout)
        slide.shapes.title.text = title
        holder = slide.placeholders[1]
        if first:
            holder.text = bullets[0] if bullets else ""
            first = False
            continue
        tf = holder.text_frame
        tf.text = bullets[0] if bullets else ""
        for b in bullets[1:]:
            para = tf.add_paragraph()
            para.text = b
            para.level = 1 if b.startswith("  ") else 0
            para.font.size = PPt(16)
    prs.core_properties.created = prs.core_properties.modified = _fixed_datetime()
    p = path(*rel.split("/"))
    prs.save(p)
    _pin_zip_times(p)
    return p


FONT_CANDIDATES = (
    "/System/Library/Fonts/Supplemental/Arial.ttf",            # macOS
    "/System/Library/Fonts/Helvetica.ttc",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",         # Debian / Ubuntu
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",                  # Fedora
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "C:/Windows/Fonts/arial.ttf",
)


def _font(size):
    """A scalable font at `size` px. The OCR images are sized for the text, so a
    fallback that ignores the size (Pillow's bitmap default) would draw ~10 px text in
    cells built for 28-51 px and quietly make the OCR test harder."""
    for candidate in FONT_CANDIDATES:
        if os.path.exists(candidate):
            try:
                return ImageFont.truetype(candidate, size)
            except OSError:
                continue
    try:
        return ImageFont.load_default(size=size)   # Pillow >= 10.1 ships a scalable one
    except TypeError:
        raise RuntimeError("no scalable font found: install DejaVu or Liberation fonts, "
                           "or Pillow >= 10.1")


def render_table_image(rel, header, rows, rotate=0, jpeg_quality=None, scale=1):
    """Render a table as an image — the screenshot-exfiltration payload for OCR tests."""
    fs = int(17 * scale)
    font, bold = _font(fs), _font(fs)
    pad, row_h = int(14 * scale), int(30 * scale)
    widths = [max(len(str(r[i])) for r in [header] + rows) for i in range(len(header))]
    col_w = [int((w * fs * 0.62) + pad * 2) for w in widths]
    W, H = sum(col_w) + pad * 2, row_h * (len(rows) + 1) + pad * 3

    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    y = pad
    x = pad
    for i, h in enumerate(header):
        d.rectangle([x, y, x + col_w[i], y + row_h], fill="#e8eaf0")
        d.text((x + pad, y + 6 * scale), str(h), font=bold, fill="black")
        x += col_w[i]
    y += row_h
    for row in rows:
        x = pad
        for i, v in enumerate(row):
            d.rectangle([x, y, x + col_w[i], y + row_h], outline="#c9ccd4")
            d.text((x + pad, y + 6 * scale), str(v), font=font, fill="black")
            x += col_w[i]
        y += row_h
    if rotate:
        img = img.rotate(rotate, expand=True, fillcolor="white")
    p = path(*rel.split("/"))
    if jpeg_quality:
        img.save(p, "JPEG", quality=jpeg_quality)
    else:
        img.save(p, "PNG")
    return p


def make_zip(rel, members, password=None):
    """Zip `members` = [(arcname, source_path), ...]. A password forces ZipCrypto.

    The archive is built at a temporary path and moved into place only once it is
    complete, so a failure never leaves the committed archive deleted."""
    p = path(*rel.split("/"))
    tmp_out = p + ".tmp"
    if password:
        # ZipCrypto via the system `zip` binary - Python's zipfile cannot write encrypted
        # archives, and ZipCrypto is what a user would actually produce from Finder/7-Zip.
        zip_bin = shutil.which("zip")
        if not zip_bin:
            raise RuntimeError(
                "password-protected zips need the `zip` command (Info-ZIP): preinstalled "
                "on macOS, `apt install zip` on Debian/Ubuntu; on Windows run the build "
                "under WSL or Git Bash with zip installed")
        staging = tempfile.mkdtemp()
        try:
            for arc, src in members:
                dst = os.path.join(staging, arc)
                os.makedirs(os.path.dirname(dst), exist_ok=True)
                shutil.copy(src, dst)
            if os.path.exists(tmp_out):
                os.remove(tmp_out)   # `zip` updates an existing archive rather than replacing it
            subprocess.run([zip_bin, "-r", "-q", "-P", password, tmp_out, "."],
                           cwd=staging, check=True)
        finally:
            shutil.rmtree(staging, ignore_errors=True)
    else:
        with zipfile.ZipFile(tmp_out, "w") as z:
            for arc, src in members:
                with open(src, "rb") as fh:
                    z.writestr(zipfile.ZipInfo(arc, date_time=FIXED_TIME), fh.read(),
                               compress_type=zipfile.ZIP_DEFLATED)
    os.replace(tmp_out, p)
    return p


def make_7z(rel, members):
    import py7zr  # imported here so only the archive stage needs it
    p = path(*rel.split("/"))
    staging = tempfile.mkdtemp()
    try:
        with py7zr.SevenZipFile(p, "w") as z:
            for arc, src in members:
                copy = os.path.join(staging, os.path.basename(arc))
                shutil.copyfile(src, copy)
                os.utime(copy, (FIXED_EPOCH, FIXED_EPOCH))  # 7z stores the file's mtime
                z.write(copy, arc)
    finally:
        shutil.rmtree(staging, ignore_errors=True)
    return p


def make_targz(rel, members):
    import gzip

    def pin(info):
        info.mtime, info.uid, info.gid, info.uname, info.gname = FIXED_EPOCH, 0, 0, "", ""
        return info

    p = path(*rel.split("/"))
    with open(p, "wb") as raw, gzip.GzipFile(fileobj=raw, mode="wb", mtime=0) as gz, \
            tarfile.open(fileobj=gz, mode="w") as t:
        for arc, src in members:
            t.add(src, arcname=arc, filter=pin)
    return p
